from datetime import datetime, timezone

from docx import Document as DOCXDocument

from packages.documents.models import Document
from packages.documents.parsers.docx import DOCXParser


def test_docx_parser_extracts_text(tmp_path):
    path = tmp_path / "contract.docx"

    source = DOCXDocument()
    source.add_paragraph("Payment is due within 30 days.")
    source.add_paragraph("The agreement expires in 2027.")
    source.save(path)

    document = Document(
        document_id="doc-docx-001",
        filename="contract.docx",
        content_hash="abc123",
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        size_bytes=path.stat().st_size,
        created_at=datetime.now(timezone.utc),
    )

    parser = DOCXParser()

    assert parser.supports(path)

    pages = parser.parse(document, path)

    assert len(pages) == 1
    assert pages[0].document_id == document.document_id
    assert pages[0].page_number == 1
    assert "Payment is due within 30 days." in pages[0].text
    assert "The agreement expires in 2027." in pages[0].text
